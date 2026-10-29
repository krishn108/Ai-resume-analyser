
import io
import re


# Job data is stored in Python; no CSV or database required.
JOBS = [
    {
        "title": "Python Developer",
        "skills": ["python", "sql", "flask", "git"]
    },
    {
        "title": "Data Analyst",
        "skills": ["python", "sql", "excel", "pandas", "data analysis"]
    },
    {
        "title": "Machine Learning Intern",
        "skills": ["python", "pandas", "numpy", "machine learning"]
    },
    {
        "title": "Web Developer",
        "skills": ["html", "css", "javascript"]
    },
    {
        "title": "Java Developer",
        "skills": ["java", "sql", "git"]
    },
    {
        "title": "Business Intelligence Intern",
        "skills": ["excel", "sql", "power bi"]
    }
]

KNOWN_SKILLS = sorted(
    {
        skill
        for job in JOBS
        for skill in job["skills"]
    },
    key=len,
    reverse=True
)


def read_resume(uploaded_file):
    """Extract text from an uploaded PDF or DOCX file."""

    if uploaded_file is None:
        raise ValueError("Please upload a resume.")

    filename = uploaded_file.name.lower()
    file_data = uploaded_file.getvalue()

    if not file_data:
        raise ValueError("The selected file is empty.")

    if filename.endswith(".pdf"):
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(file_data))

            if reader.is_encrypted:
                raise ValueError(
                    "The PDF is password-protected. Upload an unlocked PDF."
                )

            text = "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        except ValueError:
            raise
        except Exception:
            raise ValueError(
                "Unable to read this PDF. Please check the file."
            )

    elif filename.endswith(".docx"):
        from docx import Document

        try:
            document = Document(io.BytesIO(file_data))
            parts = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)

            text = "\n".join(parts)

        except Exception:
            raise ValueError(
                "Unable to read this DOCX. Please check the file."
            )

    else:
        raise ValueError("Only PDF and DOCX files are supported.")

    # Clean unnecessary spaces and blank lines.
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)
    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text found. Scanned PDFs need OCR processing."
        )

    return text


def extract_skills(resume_text):
    """Find predefined skills in the resume text."""

    text = resume_text.lower()
    found_skills = []

    for skill in KNOWN_SKILLS:
        pattern = (
            r"(?<!\w)"
            + re.escape(skill)
            + r"(?!\w)"
        )

        if re.search(pattern, text):
            found_skills.append(skill)

    return found_skills


def calculate_score(resume_text):
    """
    Score = percentage of unique predefined job-related skills found.
    This is a basic skill-coverage score.
    """

    if not KNOWN_SKILLS:
        return 0

    found = extract_skills(resume_text)

    score = round(
        len(found) * 100 / len(KNOWN_SKILLS)
    )

    return max(0, min(100, score))


def find_matching_jobs(resume_text):
    """Rank sample jobs by the percentage of required skills matched."""

    candidate_skills = extract_skills(resume_text)
    results = []

    for job in JOBS:
        required = job["skills"]

        matched = [
            skill for skill in required
            if skill in candidate_skills
        ]

        missing = [
            skill for skill in required
            if skill not in candidate_skills
        ]

        percentage = round(
            len(matched) * 100 / len(required)
        )

        results.append({
            "title": job["title"],
            "match": percentage,
            "matched": matched,
            "missing": missing
        })

    results.sort(
        key=lambda item: item["match"],
        reverse=True
    )

    return results
  
